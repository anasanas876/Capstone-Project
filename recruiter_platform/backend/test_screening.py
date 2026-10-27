import os
import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

from screening.engine import prepare_job, screen_resume
from io import BytesIO
from docx import Document

# Create a mock docx resume in memory
doc = Document()
doc.add_heading('John Doe', level=1)
doc.add_paragraph('Email: john.doe@example.com | Phone: 123-456-7890 | LinkedIn: linkedin.com/in/johndoe | GitHub: github.com/johndoe')

doc.add_heading('Skills', level=2)
doc.add_paragraph('Python, Django, React, SQL, Git, REST APIs, Problem Solving, Machine Learning')

doc.add_heading('Professional Experience', level=2)
doc.add_paragraph('Senior Software Engineer - TechCorp Inc. (City, CA) | 01-2021 - 12-2023')
doc.add_paragraph('• Developed scalable backend microservices using Python and Django.')
doc.add_paragraph('• Built REST APIs integrated with React frontends and PostgreSQL databases.')
doc.add_paragraph('• Collaborated with product managers to deliver features on schedule.')

doc.add_paragraph('Software Developer - StartupCo (City, NY) | 01-2019 - 12-2020')
doc.add_paragraph('• Engineered web applications and automated deployment pipelines using Git and Azure.')

buffer = BytesIO()
doc.save(buffer)
buffer.seek(0)

# Sample JD with Must-Have and Required Years
sample_jd = """
Senior Full-Stack Python Developer

Experience:
We are looking for an experienced developer who has architected and built robust backend systems and REST APIs using Python and Django, with modern React frontends.

Skills:
Required Skills:
- Python
- Django
- React
- SQL
- Git
- AWS
"""

job_profile = prepare_job(
    jd_text=sample_jd,
    target_skills=["python", "django", "react", "sql", "git", "aws", "docker"],
    must_have_skills=["python", "django"],
    required_experience_years=3.0
)

result = screen_resume(buffer, job_profile, filename="john_doe_resume.docx")

print("--- SCREENING RESULT ---")
print("Candidate Name:", result["candidate"]["name"])
print("Candidate Email:", result["candidate"]["email"])
print("Total Score:", result["total_score"], "/", result["max_scores"]["total"])
print("Skills Score:", result["skills_score"], "/", result["max_scores"]["skills"])
print("Experience Score:", result["experience_score"], "/", result["max_scores"]["experience"])
print("Keywords Score:", result["keywords_score"], "/", result["max_scores"]["keywords"])
print("Years of Experience:", result["years_of_experience"], "years (Required:", result["required_experience"], "years)")
print("Matched Skills:", result["matched_skills"])
print("Missing Skills:", result["missing_skills"])
print("Missing Must-Have:", result["missing_must_have"])
print("Warnings:", result["warnings"])
print("Test completed successfully!")
